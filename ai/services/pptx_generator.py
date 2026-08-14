from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_BG = RGBColor(6, 59, 57)
ACCENT = RGBColor(103, 211, 105)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(214, 234, 229)


def clean_text(value):
    return str(value or "").strip()


def add_title(slide, title):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(11.0), Inches(0.8))
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.name = "Jost"
    paragraph.font.size = Pt(32)
    paragraph.font.bold = True
    paragraph.font.color.rgb = WHITE


def add_section_badge(slide, index, time_range):
    box = slide.shapes.add_textbox(
        Inches(0.75), Inches(1.22), Inches(3.0), Inches(0.34)
    )
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = f"Section {index}"
    if time_range:
        paragraph.text += f" · {time_range}"
    paragraph.font.name = "Jost"
    paragraph.font.size = Pt(11)
    paragraph.font.bold = True
    paragraph.font.color.rgb = ACCENT


def add_body(slide, content):
    box = slide.shapes.add_textbox(Inches(0.85), Inches(1.8), Inches(10.5), Inches(4.5))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()

    lines = [line.strip() for line in clean_text(content).splitlines() if line.strip()]
    if not lines:
        lines = ["Contenu a presenter."]

    for index, line in enumerate(lines[:8]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.level = 0
        paragraph.space_after = Pt(9)
        paragraph.font.name = "Jost"
        paragraph.font.size = Pt(20 if len(lines) <= 3 else 16)
        paragraph.font.color.rgb = MUTED


def add_footer(slide):
    box = slide.shapes.add_textbox(
        Inches(0.75), Inches(6.65), Inches(11.0), Inches(0.28)
    )
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = "IWant · Speech généré"
    paragraph.alignment = PP_ALIGN.RIGHT
    paragraph.font.name = "Jost"
    paragraph.font.size = Pt(9)
    paragraph.font.color.rgb = MUTED


def apply_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = SLIDE_BG


def build_speech_pptx(speech):
    sections = speech.get("sections", []) if isinstance(speech, dict) else []
    sections = [section for section in sections if isinstance(section, dict)]
    if not sections:
        sections = [{"title": "Speech", "content": "Contenu a presenter."}]
    sections = sections[:6]

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]

    for index, section in enumerate(sections, start=1):
        slide = presentation.slides.add_slide(blank_layout)
        apply_background(slide)
        title = clean_text(section.get("title")) or f"Slide {index}"
        add_title(slide, f"{index}. {title}")
        add_section_badge(slide, index, clean_text(section.get("time_range")))
        add_body(slide, section.get("content"))
        add_footer(slide)

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
